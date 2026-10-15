"""transcript.json -> txt / md / srt / docx.

Word-for-word, one line per segment: `[HH:MM:SS] Name: text`.
Crosstalk is marked inline so readers know two people spoke at once.
"""
import os


def hms(t):
    t = int(t)
    return f"{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}"


def srt_time(t):
    ms = int(round(t * 1000))
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d},{ms % 1000:03d}"


def _crosstalk_note(seg, by_id):
    if "crosstalk" not in seg.get("flags", []):
        return ""
    others = sorted({by_id[o]["speaker"] for o in seg.get("overlaps_with", []) if o in by_id
                     and by_id[o]["speaker"] != seg["speaker"]})
    return f" [crosstalk with {', '.join(others)}]" if others else " [crosstalk]"


def _uncertain(seg):
    f = seg.get("flags", [])
    return ("low_speaker_conf" in f or "notes_mismatch" in f) and not seg.get("reviewed")


def lines(transcript, mark_uncertain=True):
    segs = sorted(transcript["segments"], key=lambda s: (s["start"], s["id"]))
    by_id = {s["id"]: s for s in segs}
    out = []
    for s in segs:
        if not s.get("text", "").strip():
            continue
        name = s["speaker"] + ("(?)" if mark_uncertain and _uncertain(s) else "")
        out.append((s, f"[{hms(s['start'])}] {name}: {s['text'].strip()}{_crosstalk_note(s, by_id)}"))
    return out


def to_txt(transcript, path):
    with open(path, "w", encoding="utf-8") as f:
        for _, ln in lines(transcript):
            f.write(ln + "\n")


def to_md(transcript, path):
    title = transcript.get("meeting", "Meeting")
    segs = transcript["segments"]
    speakers = sorted({s["speaker"] for s in segs if s.get("text")})
    n_flag = sum(1 for s in segs if _uncertain(s))
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# {title} — transcript\n\n")
        f.write(f"Speakers: {', '.join(speakers)}\n\n")
        f.write("`Name(?)` = speaker not yet confirmed by a reviewer. "
                f"{n_flag} such lines.\n\n")
        part = None
        for s, ln in lines(transcript):
            if s.get("part") != part:
                part = s.get("part")
                f.write(f"\n## Part {part + 1 if part is not None else ''} (from {hms(s['start'])})\n\n")
            f.write(ln.replace("*", r"\*") + "  \n")


def to_srt(transcript, path):
    with open(path, "w", encoding="utf-8") as f:
        n = 0
        for s, _ in lines(transcript):
            n += 1
            f.write(f"{n}\n{srt_time(s['start'])} --> {srt_time(max(s['end'], s['start'] + 0.5))}\n"
                    f"{s['speaker']}: {s['text'].strip()}\n\n")


def to_docx(transcript, path):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(f"{transcript.get('meeting', 'Meeting')} — transcript", 0)
    by_id = {s["id"]: s for s in transcript["segments"]}
    part = None
    for s, _ in lines(transcript):
        if s.get("part") != part:
            part = s.get("part")
            doc.add_heading(f"Part {part + 1 if part is not None else ''} (from {hms(s['start'])})", 2)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"[{hms(s['start'])}] ")
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        r = p.add_run(s["speaker"] + ("(?)" if _uncertain(s) else "") + ": ")
        r.bold = True
        p.add_run(s["text"].strip())
        note = _crosstalk_note(s, by_id)
        if note:
            r = p.add_run(note)
            r.italic = True
            r.font.color.rgb = RGBColor(0xB0, 0x40, 0x00)
    doc.save(path)


def export_all(transcript, out_dir, stem="transcript"):
    os.makedirs(out_dir, exist_ok=True)
    paths = {}
    for ext, fn in (("txt", to_txt), ("md", to_md), ("srt", to_srt), ("docx", to_docx)):
        p = os.path.join(out_dir, f"{stem}.{ext}")
        fn(transcript, p)
        paths[ext] = p
    return paths
